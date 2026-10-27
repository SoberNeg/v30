"""CI: проверка функций V2.0 так, как ими пользуется человек (окно Tk, Word/LibreOffice).
  • сортировка списка изделий по заголовкам;
  • имя файла «N шт. Артикул Номер Проект»;
  • паспорт без даты продажи (после подтверждения);
  • рукописная дата (опция «Печать, подпись и дата») — места, цвет, талон без печатной даты;
  • PDF-макет: автопоиск, окно обрезки (рамка мышью), PNG 300 dpi, запоминание обрезки;
  • Ctrl+C/V/X/A при русской раскладке (обработчик по символу и по коду клавиши).
Печать и подписи для теста рисуются здесь же (настоящие в репозиторий не кладутся).
Аргументы: <папка для снимков и отчёта>. Код возврата 0 — все пункты OK."""
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

src = Path(__file__).resolve().parents[1]
out = Path(sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp())
out.mkdir(parents=True, exist_ok=True)
work = Path(tempfile.mkdtemp()) / "Паспорта тест"
shutil.copytree(src, work, ignore=shutil.ignore_patterns("ci", "__pycache__", "*.паспорта.json",
                                                         "настройки.json", "Журнал_паспортов.csv"))
from PIL import Image, ImageDraw  # noqa: E402

# ---- тестовые печать и подписи (рисуем), рукописный шрифт — если есть в «Шрифты»
sd = work / "Печати и подписи"
sd.mkdir(exist_ok=True)
st = Image.new("RGBA", (600, 600), (0, 0, 0, 0))
ImageDraw.Draw(st).ellipse((10, 10, 590, 590), outline=(30, 110, 230, 255), width=14)
st.save(sd / "Печать.png")
for name, col in (("Дудин", (44, 39, 145)), ("Заливакин", (85, 80, 197))):
    im = Image.new("RGBA", (700, 420), (0, 0, 0, 0))
    ImageDraw.Draw(im).line([(20, 300), (200, 80), (330, 330), (520, 60), (680, 260)], fill=col + (255,), width=12)
    im.save(sd / f"Подпись_{name}.png")

# ---- PDF-макет: 2 страницы, синий прямоугольник в белых полях
proj_dir = work / "Проект PDF"
(proj_dir / "Макеты").mkdir(parents=True)
p1 = Image.new("RGB", (842, 595), "white")
ImageDraw.Draw(p1).rectangle((100, 150, 700, 450), fill=(30, 60, 200))
p2 = Image.new("RGB", (842, 595), "white")
ImageDraw.Draw(p2).rectangle((50, 50, 400, 300), fill=(200, 40, 40))
p1.save(proj_dir / "Макеты" / "М-0457.pdf", save_all=True, append_images=[p2], resolution=72)
shutil.copy(work / "Демо" / "Демо_проект_Одинцово.json", proj_dir / "Проект.json")

os.chdir(work)
sys.path.insert(0, str(work))
import tkinter as tk  # noqa: E402
from tkinter import messagebox  # noqa: E402

import app  # noqa: E402
import docx  # noqa: E402
import engine  # noqa: E402
import pdf_maket  # noqa: E402
import widgets  # noqa: E402

msgs, asks = [], []
for n in ("showinfo", "showwarning", "showerror"):
    setattr(messagebox, n, lambda *a, n=n, **k: msgs.append((n, a[1] if len(a) > 1 else "")))
messagebox.askyesno = lambda *a, **k: (asks.append(a[1] if len(a) > 1 else "") or True)
messagebox.askyesnocancel = lambda *a, **k: True
app.open_path = lambda p: None
results = []


def check(name, ok, detail=""):
    results.append(("OK " if ok else "ERR") + f" {name}" + (f" — {detail}" if detail else ""))


def shot(name, win):
    try:
        win.update()
        if sys.platform.startswith("win"):
            from PIL import ImageGrab
            win.update()
            x, y = win.winfo_rootx(), win.winfo_rooty()
            ImageGrab.grab((x, y, x + win.winfo_width(), y + win.winfo_height())).save(out / name)
        else:
            subprocess.run(["import", "-window", "root", str(out / name)], capture_output=True)
    except Exception as e:
        results.append(f"--- снимок {name} не сделан: {e}")


A = app.App()
A.update()
A.out_dir.set(str(work / "out"))
A.open_after.set(False)
A.make_pdf.set(False)
A.print_ready.set(False)
A._open_project(str(proj_dir / "Проект.json"))
A.update()
L = A.items_list

# 1. сортировка по заголовкам
try:
    def click(col):
        x = 0
        for (k, *_), cw in zip(L.columns, L._widths(max(L.body.winfo_width(), 200))):
            if k == col:
                L._head_click(type("E", (), {"x": x + 5})())
                return
            x += cw
    L.select(3)
    A.update()
    sel_item = L.selected
    click("_артикул")
    asc = [L.rows[i]["_артикул"] for i in L.order]
    click("_артикул")
    desc = [L.rows[i]["_артикул"] for i in L.order]
    click("_артикул")
    back = list(L.order)
    click("_размер")
    sizes = [L.rows[i]["_wh"] for i in L.order]
    check("Сортировка: артикул ▲/▼, третий щелчок — порядок проекта",
          asc == sorted(asc, key=widgets.natural_key) and desc == asc[::-1] and back == list(range(len(L.rows))),
          " | ".join(asc[:3]))
    check("Сортировка: размер как числа", sizes == sorted(sizes), str(sizes[:3]))
    check("Сортировка: выделено то же изделие", L.selected == sel_item)
    shot("v20_sort.png", A)
    L.set_sort(None, 0)
except Exception as e:
    check("Сортировка", False, repr(e))

# 2. имя файла
try:
    idx = next(i for i, it in enumerate(A.project["изделия"]) if it.get("номер_макета") == "М-0457")
    L.select(idx)
    A.update()
    name = A.name_preview.cget("text")
    q = A.project["изделия"][idx].get("тираж", 1)
    check("Имя файла «N шт. Артикул Номер Проект»", name.startswith(f"{q} шт. ") and "М-0457" in name
          and name.endswith(f'{A.project["проект"]}.docx'), name)
except Exception as e:
    check("Имя файла", False, repr(e))

# 3. PDF-макет: автопоиск, окно обрезки, запоминание
try:
    pair = A.field_widgets["ИЗОБРАЖЕНИЕ_МАКЕТА"][1]
    check("PDF: автопоиск по номеру макета", pdf_maket.is_pdf(pair.var(0, 1).get()),
          pair.sets[0]["disp"][0].get())

    def act():
        top = [w for w in A.winfo_children() if isinstance(w, tk.Toplevel)][0]
        cv = [w for w in top.winfo_children() if isinstance(w, tk.Canvas)][0]
        top.update()
        cx, cy = cv.winfo_width() // 2, cv.winfo_height() // 2
        cv.event_generate("<ButtonPress-1>", x=cx - 180, y=cy - 90)
        cv.event_generate("<B1-Motion>", x=cx + 160, y=cy + 80)
        cv.event_generate("<ButtonRelease-1>", x=cx + 160, y=cy + 80)
        top.update()
        shot("v20_crop.png", top)
        for f in top.winfo_children():
            for b in f.winfo_children():
                if b.winfo_class() == "TButton" and b.cget("text") == "Готово":
                    b.invoke()
                    return
    A.after(1200, act)
    A._crop_maket(pair, 0, 1)
    A.update()
    spec = pair.var(0, 1).get()
    _, page, box = pdf_maket.parse_spec(spec)
    check("PDF: обрезка рамкой мышью", box is not None and page == 1, pair.sets[0]["disp"][0].get())
    saved = A.store.get(A.current_item).get("макеты") or []
    check("PDF: обрезка запомнена в «<проект>.паспорта.json»",
          bool(saved) and saved[0].get("картинки", [""])[0] == spec)
    png = pdf_maket.render_png(spec, A._maket_cache_dir())
    im = Image.open(png)
    dpi = round(im.info.get("dpi", (0, 0))[0])
    check("PDF: PNG 300 dpi в «<проект>.макеты»", dpi == 300 and Path(png).parent.name.endswith(".макеты"),
          f"{im.size}, {dpi} dpi")
except Exception as e:
    check("PDF-макет", False, repr(e))

# 3б. несколько макетов на изделие: деление тиража и отдельный паспорт на каждый
try:
    mk, imgs = A._makets(), A._images()
    total = A._qty()
    mk.add_after(mk.rows[0])
    A.update()
    check("Макеты: «+» делит тираж, сумма сходится",
          mk.count() == 2 and mk.ok() and mk.distributed() == total, mk.get().replace("\n", " / "))
    mk.rows[1]["num"].set("М-0462")
    mk.rows[0]["qty"].set(total - 2)
    mk.rows[1]["qty"].set(2)
    A.update()
    caps = [st["w"][0][0].cget("text") for st in imgs.sets]
    check("Макеты: свои поля картинок у каждого", len(imgs.sets) == 2 and "М-0462" in caps[1], str(caps))
    mk.rows[1]["qty"].set(5)
    A.update()
    bad = "распределено" in mk.hint.cget("text")
    A.name_override = ""
    A._generate()
    A.update()
    check("Макеты: при несходящемся тираже паспорт не создаётся",
          bad and any("Распределено" in m[1] for m in msgs), mk.hint.cget("text"))
    mk.rows[1]["qty"].set(2)
    A.update()
    n0 = len(list((work / "out").glob("*.docx"))) if (work / "out").is_dir() else 0
    A._generate()
    A.update()
    files = sorted(p.name for p in (work / "out").glob("*.docx"))
    want = {f"{total - 2} шт.", "2 шт."}
    check("Макеты: отдельный файл на каждый макет со своим тиражом",
          len(files) - n0 == 2 and all(any(f.startswith(w) for f in files) for w in want)
          and any("М-0462" in f for f in files), " | ".join(files[-2:]))
    saved = A.store.get(A.current_item).get("макеты") or []
    check("Макеты: номера, тираж и картинки запомнены",
          len(saved) == 2 and saved[1].get("номер") == "М-0462" and saved[1].get("тираж") == 2)
    A.items_list.select(0)
    A.update()
    A.items_list.select(idx)
    A.update()
    check("Макеты: восстановлены при повторном выборе изделия",
          A._makets().get().replace("\n", " / ") == f"М-0457|{total - 2} / М-0462|2",
          A._makets().get().replace("\n", " / "))
    mk = A._makets()
    mk.remove(mk.rows[1])
    A.update()
    check("Макеты: «−» возвращает тираж", mk.count() == 1 and mk.items()[0]["тираж"] == total)
except Exception as e:
    check("Несколько макетов", False, repr(e))

# 4. паспорт без даты продажи
try:
    A.field_widgets["ДАТА_ПРОДАЖИ"][1].set("")
    asks.clear()
    A.name_override = "Тест без даты"
    A._generate()
    A.update()
    f = work / "out" / "Тест без даты.docx"
    txt = "\n".join(p.text for p in docx.Document(f).paragraphs) if f.exists() else ""
    check("Без даты: вопрос и пустое место «___» ____ 20___ г.",
          any("Дата продажи не указана" in a for a in asks) and app.BLANK_DATE in txt)
except Exception as e:
    check("Паспорт без даты", False, repr(e))

# 5. рукописная дата
try:
    A.field_widgets["ДАТА_ПРОДАЖИ"][1].set("27.10.2026")
    A.use_stamps.set(True)
    A.print_ready.set(True)
    A.name_override = "Тест рукописная дата"
    has_font = bool(app.handwriting_font())
    A._generate()
    A.update()
    f = work / "out" / "Тест рукописная дата.docx"
    d = docx.Document(f)
    descr = [dp.get("descr") for dp in d.element.iter(docx.oxml.ns.qn("wp:docPr"))]
    txt = "\n".join(p.text for p in d.paragraphs)
    if has_font:
        check("Рукописная дата: лист выдачи и талон", "Дата дудин" in descr and "Дата заливакин" in descr,
              str([x for x in descr if x and x.startswith("Дата")]))
        check("Рукописная дата: в талоне нет печатной даты", "октября 2026" not in txt)
        check("Цвет даты — из подписи", engine.ink_color(sd / "Подпись_Дудин.png")[2] > 100,
              str(engine.ink_color(sd / "Подпись_Дудин.png")))
    else:
        check("Рукописная дата: нет шрифта Caveat в «Шрифты» — место убрано, печатная дата осталась",
              "октября 2026" in txt and not any(x and x.startswith("Дата") for x in descr))
    pdf = f.with_suffix(".pdf")
    if pdf.exists() and pdf_maket.available():
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        for i in range(len(doc)):
            t = doc[i].get_textpage().get_text_range()
            if "Лист выдачи" in t[:600] or "Гарантийный талон" in t[:600]:
                doc[i].render(scale=1.6).to_pil().save(out / f"v20_page_{i + 1}.png")
        doc.close()
except Exception as e:
    check("Рукописная дата", False, repr(e))

# 6. копирование / вставка при русской раскладке
try:
    plat = widgets._platform()
    tw = tk.Toplevel(A)
    t = tk.Entry(tw, width=40)
    t.pack(padx=20, pady=20)
    tw.update()
    t.focus_force()
    A.update()
    A.clipboard_clear()
    A.clipboard_append("строка1\nстрока2\tконец")
    ev = type("E", (), {})()
    ev.widget, ev.char = t, ""
    codes = {"win": (86, 65, 67, 88), "x11": (55, 38, 54, 53), "mac": (9 << 16, 0, 8 << 16, 7 << 16)}[plat]
    ev.keysym, ev.keycode = "Cyrillic_em", codes[0]
    widgets._on_shortcut(ev)
    pasted = t.get()
    ev.keysym, ev.keycode = "Cyrillic_ef", codes[1]
    widgets._on_shortcut(ev)
    sel = t.get()[t.index("sel.first"):t.index("sel.last")] if t.selection_present() else ""
    ev.keysym, ev.keycode = "??", codes[3]                 # другая раскладка: только код клавиши
    widgets._on_shortcut(ev)
    check("Ctrl+М/Ф/Ч (русская раскладка): вставка, выделить всё, вырезать",
          pasted == "строка1 строка2 конец" and sel == pasted and t.get() == "" and
          A.clipboard_get() == pasted, f"{plat}: {pasted!r}")
    check("Коды клавиш C/V/X/A", [widgets.key_letter("", c, "", "win") for c in (67, 86, 88, 65)] == list("cvxa")
          and [widgets.key_letter("", c, "", "x11") for c in (54, 55, 53, 38)] == list("cvxa")
          and widgets.key_letter("", (8 << 16) | 0x441, "", "mac") == "c")
    ev.keysym, ev.keycode = "c", codes[2]
    check("Латиница: обработчик не дублирует стандартную привязку", widgets._on_shortcut(ev) is None)
    tw.destroy()
except Exception as e:
    check("Копирование / вставка", False, repr(e))

shot("v20_main.png", A)
A._on_close()
for m in msgs:
    results.append(f"--- {m[0]}: {m[1][:200]}".replace("\n", " | "))
print("\n".join(results))
(out / "feature_report.txt").write_text("\n".join(results), "utf-8")
sys.exit(0 if not any(r.startswith("ERR") for r in results) else 1)
